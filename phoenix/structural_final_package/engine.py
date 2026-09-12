#!/usr/bin/env python3
from __future__ import annotations

import argparse
import csv
import datetime as dt
import hashlib
import json
import shutil
import zipfile
from pathlib import Path
from typing import Any, Dict, List

STATUS = "PASS_FINAL_PRELIMINARY_STRUCTURAL_PACKAGE_ASSEMBLED_RELEASE_HOLD"
NEXT_STAGE = "PHOENIX_4.41_EXTERNAL_RELEASE_EVIDENCE_INGESTION_AND_FINAL_APPROVAL_GATE"
ENGINE_VERSION = "1.0.0"


def sha256_file(path: Path) -> str:
    h = hashlib.sha256()
    with path.open("rb") as f:
        for chunk in iter(lambda: f.read(1024 * 1024), b""):
            h.update(chunk)
    return h.hexdigest()


def load_json(path: Path) -> Dict[str, Any]:
    return json.loads(path.read_text(encoding="utf-8"))


def validate_chain(design: Dict[str,Any], solver: Dict[str,Any], report3d: Dict[str,Any],
                   qa_summary: Dict[str,Any], source_sha: str) -> None:
    if design.get("status") != "PASS_PRELIMINARY_STRUCTURAL_DESIGN_CONSOLIDATION_AND_DRAWINGS":
        raise RuntimeError("Design evidence is not PASS")
    if solver.get("status") not in (
        "PASS_REAL_OPEN_SOURCE_SOLVER_EXECUTED_PRELIMINARY_ELEMENT_VERIFICATION",
        "PASS_PRIMARY_SOLVER_EXECUTED_PRELIMINARY_ELEMENT_VERIFICATION",
    ):
        raise RuntimeError("Solver evidence is not PASS")
    if report3d.get("status") != "PASS_PRELIMINARY_3D_STRUCTURAL_MODEL_AND_CALCULATION_REPORT":
        raise RuntimeError("3D/report evidence is not PASS")
    if qa_summary.get("status") != "PASS_STRUCTURAL_QA_GATE_EXECUTED_BIM_HANDOFF_CREATED":
        raise RuntimeError("QA/BIM evidence is not PASS")

    for name, sha in [
        ("design",design.get("source_sha256")),
        ("solver",solver.get("source_sha256")),
        ("3d_report",report3d.get("source_sha256")),
        ("qa_bim",qa_summary.get("source_sha256")),
    ]:
        if sha != source_sha:
            raise RuntimeError(f"{name} source SHA mismatch")

    if qa_summary.get("release_decision") != "HOLD":
        raise RuntimeError("This stage expects current QA release decision HOLD")
    if qa_summary.get("release_allowed") is not False:
        raise RuntimeError("Release must remain locked at entry")


def closure_owner(hold_id: str) -> str:
    return {
        "H01":"Jurisdiction / competent authority + structural engineer",
        "H02":"Geotechnical engineer / site investigation provider",
        "H03":"PHOENIX after authoritative structural topology evidence",
        "H04":"Timber supplier / site verification / structural engineer",
        "H05":"PHOENIX after H01, H03 and H04 inputs",
        "H06":"PHOENIX after H01, H02 and H03 inputs",
        "H07":"Project/site survey + PHOENIX structural design",
        "H08":"Qualified structural professional",
    }.get(hold_id,"Project team")


def required_evidence(hold_id: str) -> List[str]:
    return {
        "H01":[
            "Confirmed legally applicable structural code hierarchy for Paramaribo/Suriname",
            "Confirmed governing wind design basis and basic wind parameter",
            "Confirmed consequence/reliability class and applicable partial factors",
        ],
        "H02":[
            "Site-specific geotechnical investigation (CPT/borehole or accepted local equivalent)",
            "Groundwater observation/design level",
            "Bearing/settlement parameters and foundation recommendations",
        ],
        "H03":[
            "Authoritative loadbearing wall and column topology with coordinates/axes",
            "Confirmed roof support lines and actual unsupported spans",
            "Confirmed beam/ring-beam connectivity and load path",
        ],
        "H04":[
            "Timber species",
            "Strength grade",
            "Actual dressed section dimensions",
            "Moisture/service condition and connection material data",
        ],
        "H05":[
            "Inputs H01/H03/H04 closed",
            "Connection layout and fastener/anchor specification",
            "Wind uplift design actions and complete roof anchorage load path",
        ],
        "H06":[
            "Inputs H01/H02/H03 closed",
            "Final RC material strengths and cover",
            "Full N-M/shear/anchorage/punching/detailing design",
        ],
        "H07":[
            "Tank plan location",
            "Tank support height and frame geometry",
            "Base/anchor arrangement",
            "Confirmed tank/equipment dead load and wind exposure",
        ],
        "H08":[
            "Complete final engineering package",
            "Review comments resolved",
            "Signed professional structural approval/release",
        ],
    }.get(hold_id,["Authoritative evidence sufficient to close the hold"])


def closure_plan(qa_gate: Dict[str,Any]) -> Dict[str,Any]:
    holds=[]
    auto_closed=0

    for h in qa_gate.get("release_holds",[]):
        hid=h["id"]
        prior_status=h.get("status","OPEN")
        if prior_status != "OPEN":
            status="CLOSED_FROM_PRIOR_AUTHORITY"
            auto_closed += 1
        else:
            # Current evidence intentionally cannot manufacture legal, site,
            # product, topology, geometry, or professional approval.
            status="OPEN_EXTERNAL_OR_AUTHORITY_INPUT_REQUIRED"

        holds.append({
            "id":hid,
            "hold":h["hold"],
            "prior_closure_class":h.get("closure_class",""),
            "closure_status":status,
            "responsible_party":closure_owner(hid),
            "required_evidence":required_evidence(hid),
            "phoenix_follow_up": {
                "H01":"Recalculate governed combinations/wind checks after confirmed basis is ingested.",
                "H02":"Re-run bearing, settlement/foundation sizing and RC foundation detailing.",
                "H03":"Generate authoritative structural topology and re-run global/member load paths.",
                "H04":"Re-run timber member and connection checks with confirmed material properties.",
                "H05":"Complete roof connection/uplift/anchorage design and detail drawings.",
                "H06":"Complete RC N-M/shear/punching/anchorage and reinforcement schedules.",
                "H07":"Complete tank frame, wind, overturning, anchorage and foundation design.",
                "H08":"Ingest professional review comments; release only after signed approval evidence.",
            }.get(hid,"Re-run affected checks after authoritative evidence is ingested."),
        })

    open_holds=sum(1 for h in holds if h["closure_status"].startswith("OPEN"))
    return {
        "schema":"PHOENIX_RELEASE_HOLD_CLOSURE_PLAN_1.0",
        "autonomous_closure_policy":"NO_FABRICATED_AUTHORITY_OR_SITE_EVIDENCE",
        "autonomous_holds_closed":auto_closed,
        "open_release_holds":open_holds,
        "release_decision":"HOLD" if open_holds else "READY_FOR_PROFESSIONAL_RELEASE_REVIEW",
        "holds":holds,
    }


def copy_evidence(src: Path, dst: Path) -> None:
    if not src.exists():
        raise RuntimeError(f"Evidence folder missing: {src}")
    if dst.exists():
        shutil.rmtree(dst)
    shutil.copytree(src,dst)


def write_csv(path: Path, rows: List[Dict[str,Any]], fields: List[str]) -> None:
    with path.open("w",newline="",encoding="utf-8-sig") as f:
        w=csv.DictWriter(f,fieldnames=fields)
        w.writeheader()
        for row in rows:
            out={}
            for k in fields:
                v=row.get(k,"")
                if isinstance(v,(list,dict)):
                    v=json.dumps(v,ensure_ascii=False)
                out[k]=v
            w.writerow(out)


def create_bcf_optional(plan: Dict[str,Any], out: Path) -> Dict[str,Any]:
    status={"status":"NOT_ATTEMPTED","file":None}
    try:
        from bcf.v3.bcfxml import BcfXml
        bcf=BcfXml.create_new(project_name="PHOENIX Anijstraat 616 Release Holds")
        for h in plan["holds"]:
            if not h["closure_status"].startswith("OPEN"):
                continue
            description=(
                h["hold"] + "\n\nResponsible: " + h["responsible_party"] +
                "\n\nRequired evidence:\n- " + "\n- ".join(h["required_evidence"]) +
                "\n\nPHOENIX follow-up:\n" + h["phoenix_follow_up"]
            )
            bcf.add_topic(
                title=f"{h['id']} - {h['hold']}",
                description=description,
                author="PHOENIX",
                topic_type="ReleaseHold",
                topic_status="Open",
            )
        target=out/"ANJ616_release_holds.bcf"
        bcf.save(target)
        status={"status":"PASS","file":target.name}
    except Exception as exc:
        status={"status":"SKIPPED_NONBLOCKING","file":None,"reason":str(exc)}
    (out/"bcf_export_status.json").write_text(json.dumps(status,indent=2,ensure_ascii=False),encoding="utf-8")
    return status


def build_docx(summary: Dict[str,Any], plan: Dict[str,Any], path: Path) -> None:
    from docx import Document
    from docx.shared import Mm, Pt
    from docx.enum.text import WD_ALIGN_PARAGRAPH

    doc=Document()
    sec=doc.sections[0]
    sec.top_margin=Mm(18); sec.bottom_margin=Mm(18); sec.left_margin=Mm(20); sec.right_margin=Mm(20)
    doc.styles["Normal"].font.name="Arial"
    doc.styles["Normal"].font.size=Pt(9.5)

    p=doc.add_paragraph()
    p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    r=p.add_run("FINAL PRELIMINARY STRUCTURAL PACKAGE")
    r.bold=True; r.font.size=Pt(20)
    p=doc.add_paragraph()
    p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    r=p.add_run("Woonhuis Anijstraat #616 - PHOENIX 4.41")
    r.bold=True; r.font.size=Pt(14)
    p=doc.add_paragraph()
    p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    r=p.add_run("RELEASE STATUS: HOLD - NOT FOR CONSTRUCTION")
    r.bold=True; r.font.size=Pt(13)

    doc.add_heading("1. Package status",level=1)
    doc.add_paragraph(
        "The source, load derivation, executed solver evidence, preliminary element checks, structural drawings, "
        "3D coordination model, calculation report, IFC handoff and IDS QA have been assembled into one traceable package."
    )
    tbl=doc.add_table(rows=1,cols=2); tbl.style="Table Grid"
    for i,h in enumerate(["Item","Result"]): tbl.rows[0].cells[i].text=h
    for a,b in [
        ("Project ID",summary["project_id"]),
        ("Pipeline integrity","PASS"),
        ("Autonomous release holds closed",str(plan["autonomous_holds_closed"])),
        ("Open release holds",str(plan["open_release_holds"])),
        ("Construction release","HOLD / LOCKED"),
    ]:
        cells=tbl.add_row().cells; cells[0].text=a; cells[1].text=b

    doc.add_heading("2. Release hold closure matrix",level=1)
    for h in plan["holds"]:
        doc.add_heading(f"{h['id']} - {h['hold']}",level=2)
        doc.add_paragraph(f"Status: {h['closure_status']}")
        doc.add_paragraph(f"Responsible party: {h['responsible_party']}")
        doc.add_paragraph("Required evidence:")
        for x in h["required_evidence"]:
            doc.add_paragraph(x,style="List Bullet")
        doc.add_paragraph("PHOENIX follow-up: "+h["phoenix_follow_up"])

    doc.add_heading("3. Release rule",level=1)
    doc.add_paragraph(
        "Phoenix shall not set FOR_CONSTRUCTION_RELEASE to RELEASED while any mandatory release hold remains open. "
        "External legal, geotechnical, physical material/site evidence and professional approval are not inferred."
    )
    p=doc.add_paragraph()
    r=p.add_run("PRELIMINARY / NOT FOR CONSTRUCTION - FOR_CONSTRUCTION_RELEASE = LOCKED")
    r.bold=True

    for section in doc.sections:
        fp=section.footer.paragraphs[0]
        fp.alignment=WD_ALIGN_PARAGRAPH.CENTER
        run=fp.add_run("PHOENIX 4.41 | Anijstraat #616 | FINAL PRELIMINARY PACKAGE | RELEASE HOLD")
        run.font.name="Arial"; run.font.size=Pt(8)

    doc.save(path)


def build_pdf(summary: Dict[str,Any], plan: Dict[str,Any], path: Path) -> None:
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
    from reportlab.lib.units import mm
    from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, PageBreak

    styles=getSampleStyleSheet()
    styles["BodyText"].fontName="Helvetica"; styles["BodyText"].fontSize=8.5; styles["BodyText"].leading=11
    styles["Heading1"].fontName="Helvetica-Bold"; styles["Heading1"].fontSize=14
    styles["Heading2"].fontName="Helvetica-Bold"; styles["Heading2"].fontSize=10.5
    styles.add(ParagraphStyle(name="Center",parent=styles["BodyText"],alignment=1,fontSize=11,leading=14))

    def footer(canvas,doc):
        canvas.saveState()
        canvas.setFont("Helvetica",7)
        canvas.drawCentredString(A4[0]/2,8*mm,"PHOENIX 4.41 | Anijstraat #616 | FINAL PRELIMINARY PACKAGE | RELEASE HOLD")
        canvas.restoreState()

    doc=SimpleDocTemplate(str(path),pagesize=A4,leftMargin=17*mm,rightMargin=17*mm,topMargin=17*mm,bottomMargin=16*mm)
    story=[
        Spacer(1,25*mm),
        Paragraph("<b>FINAL PRELIMINARY STRUCTURAL PACKAGE</b>",styles["Title"]),
        Spacer(1,5*mm),
        Paragraph("<b>Woonhuis Anijstraat #616 - PHOENIX 4.41</b>",styles["Center"]),
        Spacer(1,8*mm),
        Paragraph("<b>RELEASE STATUS: HOLD - NOT FOR CONSTRUCTION</b>",styles["Center"]),
        PageBreak(),
    ]

    story += [Paragraph("1. Package status",styles["Heading1"])]
    story.append(Paragraph(
        "The verified preliminary structural chain has been assembled into a single traceable package. "
        "The package is complete as an evidence handover, but construction release remains locked.",
        styles["BodyText"]
    ))
    rows=[
        ["Item","Result"],
        ["Project ID",summary["project_id"]],
        ["Pipeline integrity","PASS"],
        ["Autonomous release holds closed",str(plan["autonomous_holds_closed"])],
        ["Open release holds",str(plan["open_release_holds"])],
        ["Construction release","HOLD / LOCKED"],
    ]
    t=Table(rows,colWidths=[75*mm,90*mm],repeatRows=1)
    t.setStyle(TableStyle([
        ("GRID",(0,0),(-1,-1),0.4,colors.black),
        ("BACKGROUND",(0,0),(-1,0),colors.HexColor("#E8E8E8")),
        ("FONTNAME",(0,0),(-1,0),"Helvetica-Bold"),
        ("VALIGN",(0,0),(-1,-1),"TOP"),
        ("FONTSIZE",(0,0),(-1,-1),8.5),
    ]))
    story += [Spacer(1,3*mm),t,Spacer(1,5*mm)]

    story.append(Paragraph("2. Release hold closure matrix",styles["Heading1"]))
    for h in plan["holds"]:
        story.append(Paragraph(f"{h['id']} - {h['hold']}",styles["Heading2"]))
        story.append(Paragraph(f"<b>Status:</b> {h['closure_status']}<br/><b>Responsible:</b> {h['responsible_party']}",styles["BodyText"]))
        req="<br/>".join("• "+x for x in h["required_evidence"])
        story.append(Paragraph("<b>Required evidence:</b><br/>"+req,styles["BodyText"]))
        story.append(Paragraph("<b>PHOENIX follow-up:</b> "+h["phoenix_follow_up"],styles["BodyText"]))
        story.append(Spacer(1,3*mm))

    story.append(Paragraph("3. Release rule",styles["Heading1"]))
    story.append(Paragraph(
        "Phoenix shall not set FOR_CONSTRUCTION_RELEASE to RELEASED while any mandatory release hold remains open. "
        "Legal authority, geotechnical investigation, physical material/site facts and professional approval are never fabricated.",
        styles["BodyText"]
    ))
    story.append(Spacer(1,4*mm))
    story.append(Paragraph("<b>PRELIMINARY / NOT FOR CONSTRUCTION - FOR_CONSTRUCTION_RELEASE = LOCKED</b>",styles["BodyText"]))

    doc.build(story,onFirstPage=footer,onLaterPages=footer)


def build_input_template(plan: Dict[str,Any]) -> Dict[str,Any]:
    template={
        "schema":"PHOENIX_RELEASE_HOLD_INPUT_TEMPLATE_1.0",
        "project_id":"PHX-RP-ANIJSTRAAT-616",
        "instructions":"Fill only from authoritative evidence. Do not estimate values for release.",
        "holds":{}
    }
    for h in plan["holds"]:
        template["holds"][h["id"]]={
            "evidence_received":False,
            "evidence_source":"",
            "evidence_date":"",
            "authority_or_professional":"",
            "values_or_files":{},
            "notes":"",
            "required_evidence":h["required_evidence"],
        }
    return template


def assemble_package(design_root: Path, solver_root: Path, report3d_root: Path, qa_root: Path,
                     source_pdf: Path, out: Path) -> Dict[str,Any]:
    design=load_json(design_root/"structural_design_consolidation.json")
    solver=load_json(solver_root/"solver_element_verification.json")
    report3d=load_json(report3d_root/"structural_3d_report_summary.json")
    qa_summary=load_json(qa_root/"structural_QA_BIM_summary.json")
    qa_gate=load_json(qa_root/"structural_QA_release_gate.json")
    source_sha=sha256_file(source_pdf)

    validate_chain(design,solver,report3d,qa_summary,source_sha)
    plan=closure_plan(qa_gate)

    out.mkdir(parents=True,exist_ok=True)
    evidence=out/"evidence"
    evidence.mkdir(exist_ok=True)

    copy_evidence(design_root,evidence/"01_design_drawings")
    copy_evidence(solver_root,evidence/"02_solver_verification")
    copy_evidence(report3d_root,evidence/"03_3d_calculation_report")
    copy_evidence(qa_root,evidence/"04_qa_bim")
    source_dir=evidence/"00_source"
    source_dir.mkdir()
    shutil.copy2(source_pdf,source_dir/"ANJ616_COMPLETE_DESIGN_SOURCE.pdf")

    (out/"release_hold_closure_plan.json").write_text(json.dumps(plan,indent=2,ensure_ascii=False),encoding="utf-8")
    write_csv(
        out/"release_hold_closure_matrix.csv",
        plan["holds"],
        ["id","hold","prior_closure_class","closure_status","responsible_party","required_evidence","phoenix_follow_up"]
    )

    template=build_input_template(plan)
    (out/"RELEASE_HOLD_INPUT_TEMPLATE.json").write_text(json.dumps(template,indent=2,ensure_ascii=False),encoding="utf-8")

    requests=out/"hold_closure_requests"
    requests.mkdir()
    for h in plan["holds"]:
        lines=[
            f"# {h['id']} - Release Hold Closure Request",
            "",
            f"**Hold:** {h['hold']}",
            "",
            f"**Responsible party:** {h['responsible_party']}",
            "",
            "## Required authoritative evidence",
            "",
        ]
        lines += [f"- {x}" for x in h["required_evidence"]]
        lines += [
            "",
            "## Phoenix follow-up after evidence ingestion",
            "",
            h["phoenix_follow_up"],
            "",
            "**Current status:** OPEN",
            "",
            "**Construction release remains locked.**",
        ]
        (requests/f"{h['id']}_closure_request.md").write_text("\n".join(lines),encoding="utf-8")

    summary={
        "schema":"PHOENIX_FINAL_PRELIMINARY_STRUCTURAL_PACKAGE_1.0",
        "project_id":design["project_id"],
        "source_sha256":source_sha,
        "pipeline_integrity":"PASS",
        "design_evidence":"VERIFIED",
        "solver_evidence":"VERIFIED",
        "report3d_evidence":"VERIFIED",
        "qa_bim_evidence":"VERIFIED",
        "autonomous_holds_closed":plan["autonomous_holds_closed"],
        "open_release_holds":plan["open_release_holds"],
        "release_decision":"HOLD" if plan["open_release_holds"] else "READY_FOR_PROFESSIONAL_RELEASE_REVIEW",
        "for_construction_release":"LOCKED",
        "professional_structural_review_required":True,
        "status":STATUS,
        "next_stage":NEXT_STAGE,
    }

    (out/"final_structural_package_summary.json").write_text(json.dumps(summary,indent=2),encoding="utf-8")

    bcf_status=create_bcf_optional(plan,out/"hold_closure_requests")
    summary["bcf_export"]=bcf_status["status"]
    (out/"final_structural_package_summary.json").write_text(json.dumps(summary,indent=2),encoding="utf-8")

    build_docx(summary,plan,out/"PHOENIX_ANIJSTRAAT_616_FINAL_PRELIMINARY_STRUCTURAL_PACKAGE_REPORT.docx")
    build_pdf(summary,plan,out/"PHOENIX_ANIJSTRAAT_616_FINAL_PRELIMINARY_STRUCTURAL_PACKAGE_REPORT.pdf")

    readme=[
        "# PHOENIX Anijstraat #616 - Final Preliminary Structural Package",
        "",
        "This is the consolidated evidence package after the structural QA/BIM gate.",
        "",
        f"- Pipeline integrity: PASS",
        f"- Open release holds: {plan['open_release_holds']}",
        f"- Construction release: LOCKED",
        f"- Professional review required: TRUE",
        "",
        "The package is suitable for evidence handover and release-hold closure work.",
        "It is not a for-construction release.",
    ]
    (out/"README_RELEASE_STATUS.md").write_text("\n".join(readme),encoding="utf-8")

    # Manifest before archive, excluding archive itself.
    manifest={}
    for p in sorted(out.rglob("*")):
        if p.is_file() and p.name not in ("evidence_manifest.json","PHOENIX_ANIJSTRAAT_616_FINAL_PRELIMINARY_STRUCTURAL_PACKAGE.zip"):
            manifest[p.relative_to(out).as_posix()]={"sha256":sha256_file(p),"bytes":p.stat().st_size}
    (out/"evidence_manifest.json").write_text(json.dumps(manifest,indent=2),encoding="utf-8")

    archive=out/"PHOENIX_ANIJSTRAAT_616_FINAL_PRELIMINARY_STRUCTURAL_PACKAGE.zip"
    with zipfile.ZipFile(archive,"w",zipfile.ZIP_DEFLATED) as z:
        for p in sorted(out.rglob("*")):
            if p.is_file() and p != archive:
                z.write(p,p.relative_to(out).as_posix())

    summary["final_archive"]=archive.name
    summary["final_archive_sha256"]=sha256_file(archive)
    (out/"final_structural_package_summary.json").write_text(json.dumps(summary,indent=2),encoding="utf-8")
    return summary


def self_test() -> None:
    qa_gate={"release_holds":[
        {"id":f"H{i:02d}","hold":"test","closure_class":"X","status":"OPEN"}
        for i in range(1,9)
    ]}
    plan=closure_plan(qa_gate)
    assert plan["open_release_holds"]==8
    assert plan["autonomous_holds_closed"]==0
    assert plan["release_decision"]=="HOLD"
    assert plan["holds"][-1]["responsible_party"]=="Qualified structural professional"
    print("PHOENIX_4_41_FINAL_STRUCTURAL_PACKAGE_SELF_TEST=PASS")


def verify_output(path: Path) -> None:
    d=load_json(path)
    out=path.parent
    assert d["status"]==STATUS
    assert d["pipeline_integrity"]=="PASS"
    assert d["release_decision"]=="HOLD"
    assert d["for_construction_release"]=="LOCKED"
    assert d["open_release_holds"]>=1
    required=[
        "release_hold_closure_plan.json",
        "release_hold_closure_matrix.csv",
        "RELEASE_HOLD_INPUT_TEMPLATE.json",
        "PHOENIX_ANIJSTRAAT_616_FINAL_PRELIMINARY_STRUCTURAL_PACKAGE_REPORT.docx",
        "PHOENIX_ANIJSTRAAT_616_FINAL_PRELIMINARY_STRUCTURAL_PACKAGE_REPORT.pdf",
        "README_RELEASE_STATUS.md",
        "evidence_manifest.json",
        "PHOENIX_ANIJSTRAAT_616_FINAL_PRELIMINARY_STRUCTURAL_PACKAGE.zip",
    ]
    for rel in required:
        p=out/rel
        assert p.exists() and p.stat().st_size>0, rel

    from pypdf import PdfReader
    pdf=PdfReader(str(out/"PHOENIX_ANIJSTRAAT_616_FINAL_PRELIMINARY_STRUCTURAL_PACKAGE_REPORT.pdf"))
    assert len(pdf.pages)>=2
    from docx import Document
    doc=Document(str(out/"PHOENIX_ANIJSTRAAT_616_FINAL_PRELIMINARY_STRUCTURAL_PACKAGE_REPORT.docx"))
    assert len(doc.paragraphs)>20

    with zipfile.ZipFile(out/"PHOENIX_ANIJSTRAAT_616_FINAL_PRELIMINARY_STRUCTURAL_PACKAGE.zip") as z:
        names=set(z.namelist())
        assert "final_structural_package_summary.json" in names
        assert "evidence/00_source/ANJ616_COMPLETE_DESIGN_SOURCE.pdf" in names

    print("PHOENIX_4_41_FINAL_STRUCTURAL_PACKAGE_OUTPUT_VERIFY=PASS")


def main() -> int:
    ap=argparse.ArgumentParser()
    ap.add_argument("--design-root",type=Path)
    ap.add_argument("--solver-root",type=Path)
    ap.add_argument("--report3d-root",type=Path)
    ap.add_argument("--qa-root",type=Path)
    ap.add_argument("--source-pdf",type=Path)
    ap.add_argument("--output",type=Path)
    ap.add_argument("--self-test",action="store_true")
    ap.add_argument("--verify-output",type=Path)
    args=ap.parse_args()

    if args.self_test:
        self_test(); return 0
    if args.verify_output:
        verify_output(args.verify_output); return 0
    if not all([args.design_root,args.solver_root,args.report3d_root,args.qa_root,args.source_pdf,args.output]):
        ap.error("--design-root --solver-root --report3d-root --qa-root --source-pdf --output required")

    result=assemble_package(
        args.design_root,args.solver_root,args.report3d_root,args.qa_root,args.source_pdf,args.output
    )
    print(f"PROJECT_ID={result['project_id']}")
    print("PIPELINE_INTEGRITY=PASS")
    print(f"AUTONOMOUS_RELEASE_HOLDS_CLOSED={result['autonomous_holds_closed']}")
    print(f"OPEN_RELEASE_HOLDS={result['open_release_holds']}")
    print(f"BCF_EXPORT={result['bcf_export']}")
    print("FINAL_PRELIMINARY_STRUCTURAL_PACKAGE=PASS")
    print("RELEASE_DECISION=HOLD")
    print("FOR_CONSTRUCTION_RELEASE=LOCKED")
    print(f"FINAL_PACKAGE_ARCHIVE_SHA256={result['final_archive_sha256']}")
    print(f"STRUCTURAL_FINAL_PACKAGE_STATUS={result['status']}")
    print(f"NEXT_STAGE={result['next_stage']}")
    print(f"OUTPUT={args.output}")
    return 0


if __name__=="__main__":
    raise SystemExit(main())
