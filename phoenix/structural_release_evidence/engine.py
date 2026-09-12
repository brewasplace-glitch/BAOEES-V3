#!/usr/bin/env python3
from __future__ import annotations

import argparse
import csv
import datetime as dt
import hashlib
import json
import shutil
import zipfile
from pathlib import Path
from typing import Any, Dict, List, Tuple

STATUS = "PASS_EXTERNAL_RELEASE_EVIDENCE_INGESTION_GATE_EXECUTED"
NEXT_STAGE = "PHOENIX_4.41_RELEASE_EVIDENCE_PROCESSING_LOOP_AND_SIGNED_RELEASE_ACTION"
ENGINE_VERSION = "1.0.0"


def sha256_file(path: Path) -> str:
    h=hashlib.sha256()
    with path.open("rb") as f:
        for chunk in iter(lambda:f.read(1024*1024),b""):
            h.update(chunk)
    return h.hexdigest()


def load_json(path: Path) -> Dict[str,Any]:
    return json.loads(path.read_text(encoding="utf-8"))


def hold_requirements() -> Dict[str,List[str]]:
    return {
        "H01":[
            "Confirmed legally applicable structural code hierarchy for Paramaribo/Suriname",
            "Confirmed governing wind design basis/basic wind parameter",
            "Confirmed reliability/consequence basis and applicable partial factors",
        ],
        "H02":[
            "Site-specific geotechnical investigation",
            "Groundwater observation/design level",
            "Bearing, settlement and foundation recommendation parameters",
        ],
        "H03":[
            "Authoritative loadbearing wall and column topology",
            "Confirmed roof support lines and actual unsupported spans",
            "Confirmed structural connectivity/load path",
        ],
        "H04":[
            "Timber species and strength grade",
            "Actual dressed section dimensions",
            "Moisture/service condition and connection material data",
        ],
        "H05":[
            "H01/H03/H04 governing inputs",
            "Connection and fastener/anchor specification",
            "Wind uplift actions and complete roof anchorage load path",
        ],
        "H06":[
            "H01/H02/H03 governing inputs",
            "Final concrete/reinforcement material strengths and cover",
            "N-M/shear/anchorage/punching/detailing evidence",
        ],
        "H07":[
            "Tank plan location",
            "Tank support height/frame geometry",
            "Base/anchor arrangement and confirmed equipment load",
        ],
        "H08":[
            "Qualified structural professional review",
            "Resolved review comments",
            "Signed approval evidence with clearly stated scope",
        ],
    }


def bootstrap_inbox(root: Path, schema_path: Path) -> None:
    root.mkdir(parents=True,exist_ok=True)
    reqs=hold_requirements()
    readme=[
        "# PHOENIX - Anijstraat #616 External Release Evidence Inbox",
        "",
        "Place authoritative release evidence in the corresponding H01-H08 folder.",
        "For a submission to be ingested, copy `submission.template.json` to `submission.json` and complete it.",
        "",
        "Important:",
        "- File presence does not close a release hold.",
        "- JSON schema validity does not prove legal authority or professional qualification.",
        "- Optional Ed25519 signatures verify file integrity/authenticity against the supplied key; they do not establish professional competence.",
        "- Construction release remains locked until the controlled final release action.",
    ]
    (root/"README.md").write_text("\n".join(readme),encoding="utf-8")
    shutil.copy2(schema_path,root/"release_evidence_submission.schema.json")

    for hid, required in reqs.items():
        d=root/hid
        d.mkdir(exist_ok=True)
        template={
            "schema_version":"1.0",
            "project_id":"PHX-RP-ANIJSTRAAT-616",
            "hold_id":hid,
            "issuer":{
                "name":"",
                "organization":"",
                "role":"",
                "registration_or_authority_id":"",
                "contact_reference":""
            },
            "evidence_date":"YYYY-MM-DD",
            "evidence_files":[
                {"path":"example.pdf","document_type":"authoritative evidence"}
            ],
            "assertions":required,
            "authority_basis":"",
            "notes":"",
            "professional_approval":{
                "decision":"NOT_APPLICABLE" if hid!="H08" else "REVIEW_ONLY",
                "scope":"",
                "credential_confirmation":False
            }
        }
        (d/"submission.template.json").write_text(json.dumps(template,indent=2,ensure_ascii=False),encoding="utf-8")
        (d/"REQUIRED_EVIDENCE.md").write_text(
            "# "+hid+" Required Evidence\n\n"+"\n".join("- "+x for x in required)+
            "\n\nCurrent release state: OPEN.\n",
            encoding="utf-8"
        )


def validate_prior(final_root: Path, source_sha: str) -> Tuple[Dict[str,Any],Dict[str,Any]]:
    summary=load_json(final_root/"final_structural_package_summary.json")
    plan=load_json(final_root/"release_hold_closure_plan.json")

    if summary.get("status")!="PASS_FINAL_PRELIMINARY_STRUCTURAL_PACKAGE_ASSEMBLED_RELEASE_HOLD":
        raise RuntimeError("Prior final preliminary structural package is not PASS")
    if summary.get("source_sha256")!=source_sha:
        raise RuntimeError("Prior final package source SHA mismatch")
    if summary.get("for_construction_release")!="LOCKED":
        raise RuntimeError("Expected prior release lock")
    if summary.get("release_decision")!="HOLD":
        raise RuntimeError("Expected prior release decision HOLD")
    if len(plan.get("holds",[]))!=8:
        raise RuntimeError("Expected eight release holds in prior package")
    return summary,plan


def validate_schema(instance: Dict[str,Any], schema: Dict[str,Any]) -> Tuple[bool,List[str]]:
    import jsonschema
    validator=jsonschema.Draft202012Validator(schema,format_checker=jsonschema.FormatChecker())
    errors=sorted(validator.iter_errors(instance),key=lambda e:list(e.path))
    return (len(errors)==0,[e.message for e in errors])


def verify_signature(submission: Dict[str,Any], hold_dir: Path) -> Dict[str,Any]:
    sig=submission.get("signature")
    if not sig:
        return {"status":"NOT_PROVIDED"}

    try:
        from cryptography.hazmat.primitives import serialization
        from cryptography.hazmat.primitives.asymmetric.ed25519 import Ed25519PublicKey

        if sig.get("algorithm")!="ED25519":
            return {"status":"UNSUPPORTED_ALGORITHM"}

        signed=hold_dir/sig["signed_file"]
        signature=hold_dir/sig["signature_file"]
        public_key=hold_dir/sig["public_key_pem_file"]
        for p in [signed,signature,public_key]:
            if not p.exists() or not p.is_file():
                return {"status":"MISSING_SIGNATURE_ARTIFACT","missing":str(p.name)}

        key=serialization.load_pem_public_key(public_key.read_bytes())
        if not isinstance(key,Ed25519PublicKey):
            return {"status":"INVALID_PUBLIC_KEY_TYPE"}
        key.verify(signature.read_bytes(),signed.read_bytes())
        return {
            "status":"VERIFIED",
            "signed_file":signed.name,
            "signed_file_sha256":sha256_file(signed),
        }
    except Exception as exc:
        return {"status":"INVALID","reason":str(exc)}


def ingest_hold(hold: Dict[str,Any], evidence_root: Path, schema: Dict[str,Any], snapshot_root: Path) -> Dict[str,Any]:
    hid=hold["id"]
    hold_dir=evidence_root/hid
    submission_path=hold_dir/"submission.json"

    base={
        "id":hid,
        "hold":hold["hold"],
        "prior_status":hold.get("closure_status","OPEN"),
        "submission_status":"OPEN_NO_EXTERNAL_EVIDENCE",
        "integrity_status":"NOT_EVALUATED",
        "signature_status":"NOT_PROVIDED",
        "domain_acceptance":"NOT_PERFORMED",
        "closure_status":"OPEN",
        "evidence_files":[],
        "issuer":None,
        "evidence_date":None,
    }

    if not submission_path.exists():
        return base

    try:
        submission=load_json(submission_path)
    except Exception as exc:
        base["submission_status"]="REJECTED_INVALID_JSON"
        base["integrity_status"]="FAIL"
        base["errors"]=[str(exc)]
        return base

    ok,errors=validate_schema(submission,schema)
    if not ok:
        base["submission_status"]="REJECTED_SCHEMA"
        base["integrity_status"]="FAIL"
        base["errors"]=errors
        return base

    if submission["hold_id"]!=hid:
        base["submission_status"]="REJECTED_HOLD_ID_MISMATCH"
        base["integrity_status"]="FAIL"
        return base

    evidence=[]
    integrity_ok=True
    for item in submission["evidence_files"]:
        rel=Path(item["path"])
        if rel.is_absolute() or ".." in rel.parts:
            evidence.append({"path":item["path"],"status":"INVALID_PATH"})
            integrity_ok=False
            continue
        p=hold_dir/rel
        if not p.exists() or not p.is_file():
            evidence.append({"path":item["path"],"status":"MISSING"})
            integrity_ok=False
            continue
        sha=sha256_file(p)
        expected=item.get("sha256")
        status="PASS"
        if expected and expected.lower()!=sha:
            status="HASH_MISMATCH"
            integrity_ok=False
        evidence.append({
            "path":item["path"],
            "document_type":item["document_type"],
            "sha256":sha,
            "expected_sha256":expected,
            "status":status,
            "bytes":p.stat().st_size,
        })

    sig=verify_signature(submission,hold_dir)
    if sig["status"] in ("INVALID","MISSING_SIGNATURE_ARTIFACT","INVALID_PUBLIC_KEY_TYPE","UNSUPPORTED_ALGORITHM"):
        integrity_ok=False

    if not integrity_ok:
        base.update({
            "submission_status":"REJECTED_INTEGRITY",
            "integrity_status":"FAIL",
            "signature_status":sig["status"],
            "evidence_files":evidence,
            "issuer":submission["issuer"],
            "evidence_date":submission["evidence_date"],
        })
        return base

    dest=snapshot_root/hid
    if dest.exists():
        shutil.rmtree(dest)
    shutil.copytree(hold_dir,dest)

    base.update({
        "submission_status":"INGESTED_SCHEMA_AND_INTEGRITY_PASS",
        "integrity_status":"PASS",
        "signature_status":sig["status"],
        "evidence_files":evidence,
        "issuer":submission["issuer"],
        "evidence_date":submission["evidence_date"],
        "assertions":submission["assertions"],
        "authority_basis":submission.get("authority_basis",""),
        "professional_approval":submission.get("professional_approval",{}),
        "domain_acceptance":"PENDING_AUTHORITY_ENGINEERING_OR_PROFESSIONAL_VALIDATION",
        "closure_status":"OPEN_EVIDENCE_INGESTED_PENDING_ACCEPTANCE",
    })
    return base


def approval_gate(records: List[Dict[str,Any]]) -> Dict[str,Any]:
    open_count=sum(1 for r in records if not str(r["closure_status"]).startswith("CLOSED"))
    ingested=sum(1 for r in records if r["submission_status"]=="INGESTED_SCHEMA_AND_INTEGRITY_PASS")
    rejected=sum(1 for r in records if str(r["submission_status"]).startswith("REJECTED"))

    h08=next((r for r in records if r["id"]=="H08"),None)
    professional_approval_evidence=False
    if h08 and h08["submission_status"]=="INGESTED_SCHEMA_AND_INTEGRITY_PASS":
        p=h08.get("professional_approval") or {}
        professional_approval_evidence=(
            p.get("decision")=="APPROVED_FOR_CONSTRUCTION"
            and p.get("credential_confirmation") is True
        )

    if ingested==8 and rejected==0 and professional_approval_evidence:
        state="HOLD_PENDING_DOMAIN_ACCEPTANCE_RECALCULATION_AND_CONTROLLED_RELEASE_ACTION"
    elif ingested>0:
        state="HOLD_EXTERNAL_EVIDENCE_PARTIALLY_INGESTED"
    else:
        state="HOLD_NO_EXTERNAL_RELEASE_EVIDENCE_INGESTED"

    return {
        "schema":"PHOENIX_FINAL_APPROVAL_GATE_1.0",
        "gate_state":state,
        "submission_integrity_pass_count":ingested,
        "rejected_submission_count":rejected,
        "open_release_holds":open_count,
        "professional_approval_evidence_declared":professional_approval_evidence,
        "construction_release":"LOCKED",
        "release_allowed":False,
        "reason":"This ingestion gate never converts file presence, schema validity, hashes, or signatures into engineering/legal/professional acceptance.",
    }


def write_csv(path: Path, rows: List[Dict[str,Any]]) -> None:
    fields=[
        "id","hold","prior_status","submission_status","integrity_status",
        "signature_status","domain_acceptance","closure_status","evidence_date"
    ]
    with path.open("w",newline="",encoding="utf-8-sig") as f:
        w=csv.DictWriter(f,fieldnames=fields)
        w.writeheader()
        for row in rows:
            w.writerow({k:row.get(k,"") for k in fields})


def build_docx(gate: Dict[str,Any], records: List[Dict[str,Any]], path: Path) -> None:
    from docx import Document
    from docx.shared import Mm, Pt
    from docx.enum.text import WD_ALIGN_PARAGRAPH

    doc=Document()
    sec=doc.sections[0]
    sec.top_margin=Mm(18); sec.bottom_margin=Mm(18); sec.left_margin=Mm(20); sec.right_margin=Mm(20)
    doc.styles["Normal"].font.name="Arial"; doc.styles["Normal"].font.size=Pt(9.5)

    p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    r=p.add_run("EXTERNAL RELEASE EVIDENCE INGESTION + FINAL APPROVAL GATE")
    r.bold=True; r.font.size=Pt(18)
    p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    p.add_run("Woonhuis Anijstraat #616 - PHOENIX 4.41").bold=True
    p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    r=p.add_run("FOR_CONSTRUCTION_RELEASE = LOCKED"); r.bold=True; r.font.size=Pt(12)

    doc.add_heading("1. Gate result",level=1)
    tbl=doc.add_table(rows=1,cols=2); tbl.style="Table Grid"
    tbl.rows[0].cells[0].text="Item"; tbl.rows[0].cells[1].text="Result"
    for a,b in [
        ("Gate state",gate["gate_state"]),
        ("Evidence submissions integrity PASS",str(gate["submission_integrity_pass_count"])),
        ("Rejected submissions",str(gate["rejected_submission_count"])),
        ("Open release holds",str(gate["open_release_holds"])),
        ("Construction release",gate["construction_release"]),
    ]:
        cells=tbl.add_row().cells; cells[0].text=a; cells[1].text=b

    doc.add_heading("2. Evidence intake by release hold",level=1)
    for rec in records:
        doc.add_heading(f"{rec['id']} - {rec['hold']}",level=2)
        doc.add_paragraph(f"Submission: {rec['submission_status']}")
        doc.add_paragraph(f"Integrity: {rec['integrity_status']}; signature: {rec['signature_status']}")
        doc.add_paragraph(f"Domain acceptance: {rec['domain_acceptance']}")
        doc.add_paragraph(f"Closure status: {rec['closure_status']}")
        if rec.get("issuer"):
            i=rec["issuer"]
            doc.add_paragraph(f"Issuer declared: {i.get('name','')} / {i.get('organization','')} / {i.get('role','')}")
        if rec.get("evidence_files"):
            doc.add_paragraph("Evidence files:")
            for e in rec["evidence_files"]:
                doc.add_paragraph(f"{e.get('path')} - {e.get('status')} - {e.get('sha256','')}",style="List Bullet")

    doc.add_heading("3. Governance",level=1)
    doc.add_paragraph(gate["reason"])
    doc.add_paragraph(
        "A verified digital signature proves cryptographic integrity relative to a supplied public key; "
        "it does not prove that the issuer is a competent authority or qualified structural professional."
    )
    p=doc.add_paragraph()
    r=p.add_run("PRELIMINARY / NOT FOR CONSTRUCTION - RELEASE ACTION NOT AUTHORIZED"); r.bold=True

    for s in doc.sections:
        fp=s.footer.paragraphs[0]; fp.alignment=WD_ALIGN_PARAGRAPH.CENTER
        rr=fp.add_run("PHOENIX 4.41 | Anijstraat #616 | External Evidence Gate | RELEASE LOCKED")
        rr.font.name="Arial"; rr.font.size=Pt(8)
    doc.save(path)


def build_pdf(gate: Dict[str,Any], records: List[Dict[str,Any]], path: Path) -> None:
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import getSampleStyleSheet
    from reportlab.lib.units import mm
    from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle

    styles=getSampleStyleSheet()
    styles["BodyText"].fontName="Helvetica"; styles["BodyText"].fontSize=8.5; styles["BodyText"].leading=11
    styles["Heading1"].fontName="Helvetica-Bold"; styles["Heading1"].fontSize=14
    styles["Heading2"].fontName="Helvetica-Bold"; styles["Heading2"].fontSize=10.5

    def footer(canvas,doc):
        canvas.saveState(); canvas.setFont("Helvetica",7)
        canvas.drawCentredString(A4[0]/2,8*mm,"PHOENIX 4.41 | Anijstraat #616 | External Evidence Gate | RELEASE LOCKED")
        canvas.restoreState()

    doc=SimpleDocTemplate(str(path),pagesize=A4,leftMargin=17*mm,rightMargin=17*mm,topMargin=17*mm,bottomMargin=16*mm)
    story=[
        Paragraph("EXTERNAL RELEASE EVIDENCE INGESTION + FINAL APPROVAL GATE",styles["Title"]),
        Spacer(1,3*mm),
        Paragraph("<b>Woonhuis Anijstraat #616 - PHOENIX 4.41</b>",styles["BodyText"]),
        Spacer(1,3*mm),
        Paragraph("<b>FOR_CONSTRUCTION_RELEASE = LOCKED</b>",styles["BodyText"]),
        Spacer(1,5*mm),
        Paragraph("1. Gate result",styles["Heading1"]),
    ]
    rows=[
        ["Item","Result"],
        ["Gate state",gate["gate_state"]],
        ["Evidence submissions integrity PASS",str(gate["submission_integrity_pass_count"])],
        ["Rejected submissions",str(gate["rejected_submission_count"])],
        ["Open release holds",str(gate["open_release_holds"])],
        ["Construction release",gate["construction_release"]],
    ]
    table=Table(rows,colWidths=[65*mm,100*mm],repeatRows=1)
    table.setStyle(TableStyle([
        ("GRID",(0,0),(-1,-1),0.4,colors.black),
        ("BACKGROUND",(0,0),(-1,0),colors.HexColor("#E8E8E8")),
        ("FONTNAME",(0,0),(-1,0),"Helvetica-Bold"),
        ("FONTSIZE",(0,0),(-1,-1),8),
        ("VALIGN",(0,0),(-1,-1),"TOP"),
    ]))
    story += [table,Spacer(1,5*mm),Paragraph("2. Evidence intake by release hold",styles["Heading1"])]

    for rec in records:
        story.append(Paragraph(f"{rec['id']} - {rec['hold']}",styles["Heading2"]))
        text=(
            f"Submission: {rec['submission_status']}<br/>"
            f"Integrity: {rec['integrity_status']} | Signature: {rec['signature_status']}<br/>"
            f"Domain acceptance: {rec['domain_acceptance']}<br/>"
            f"Closure status: {rec['closure_status']}"
        )
        story.append(Paragraph(text,styles["BodyText"]))
        story.append(Spacer(1,2*mm))

    story += [
        Paragraph("3. Governance",styles["Heading1"]),
        Paragraph(gate["reason"],styles["BodyText"]),
        Paragraph(
            "Digital signature verification is cryptographic integrity evidence only. It does not establish legal authority, "
            "professional registration, engineering correctness, or approval scope.",
            styles["BodyText"]
        ),
        Spacer(1,3*mm),
        Paragraph("<b>PRELIMINARY / NOT FOR CONSTRUCTION - RELEASE ACTION NOT AUTHORIZED</b>",styles["BodyText"]),
    ]
    doc.build(story,onFirstPage=footer,onLaterPages=footer)


def run_gate(final_root: Path, evidence_root: Path, schema_path: Path, source_pdf: Path, out: Path) -> Dict[str,Any]:
    source_sha=sha256_file(source_pdf)
    prior,plan=validate_prior(final_root,source_sha)
    schema=load_json(schema_path)

    out.mkdir(parents=True,exist_ok=True)
    snapshot=out/"ingested_evidence_snapshot"
    snapshot.mkdir(exist_ok=True)

    records=[ingest_hold(h,evidence_root,schema,snapshot) for h in plan["holds"]]
    gate=approval_gate(records)

    (out/"evidence_intake_register.json").write_text(json.dumps(records,indent=2,ensure_ascii=False),encoding="utf-8")
    write_csv(out/"evidence_intake_register.csv",records)
    (out/"final_approval_gate.json").write_text(json.dumps(gate,indent=2),encoding="utf-8")
    shutil.copy2(schema_path,out/"release_evidence_submission.schema.json")

    # Reusable processing instructions.
    readme=[
        "# PHOENIX Anijstraat #616 - External Release Evidence Gate",
        "",
        f"Gate state: {gate['gate_state']}",
        f"Integrity-pass submissions: {gate['submission_integrity_pass_count']}",
        f"Open release holds: {gate['open_release_holds']}",
        "",
        "Construction release remains LOCKED.",
        "",
        "To process later evidence without reinstalling the capability, rerun the repository runner with the same evidence-root folder.",
    ]
    (out/"README_GATE_STATUS.md").write_text("\n".join(readme),encoding="utf-8")

    build_docx(gate,records,out/"PHOENIX_ANIJSTRAAT_616_EXTERNAL_RELEASE_EVIDENCE_APPROVAL_GATE.docx")
    build_pdf(gate,records,out/"PHOENIX_ANIJSTRAAT_616_EXTERNAL_RELEASE_EVIDENCE_APPROVAL_GATE.pdf")

    summary={
        "schema":"PHOENIX_EXTERNAL_RELEASE_EVIDENCE_GATE_SUMMARY_1.0",
        "project_id":"PHX-RP-ANIJSTRAAT-616",
        "source_sha256":source_sha,
        "prior_final_package_status":prior["status"],
        "ingested_submission_count":gate["submission_integrity_pass_count"],
        "rejected_submission_count":gate["rejected_submission_count"],
        "open_release_holds":gate["open_release_holds"],
        "gate_state":gate["gate_state"],
        "construction_release":"LOCKED",
        "release_allowed":False,
        "status":STATUS,
        "next_stage":NEXT_STAGE,
    }

    (out/"external_release_evidence_gate_summary.json").write_text(json.dumps(summary,indent=2),encoding="utf-8")

    manifest={}
    for p in sorted(out.rglob("*")):
        if p.is_file() and p.name not in ("evidence_manifest.json","PHOENIX_ANIJSTRAAT_616_EXTERNAL_RELEASE_EVIDENCE_GATE_PACKAGE.zip"):
            manifest[p.relative_to(out).as_posix()]={"sha256":sha256_file(p),"bytes":p.stat().st_size}
    (out/"evidence_manifest.json").write_text(json.dumps(manifest,indent=2),encoding="utf-8")

    archive=out/"PHOENIX_ANIJSTRAAT_616_EXTERNAL_RELEASE_EVIDENCE_GATE_PACKAGE.zip"
    with zipfile.ZipFile(archive,"w",zipfile.ZIP_DEFLATED) as z:
        for p in sorted(out.rglob("*")):
            if p.is_file() and p!=archive:
                z.write(p,p.relative_to(out).as_posix())

    summary["archive_sha256"]=sha256_file(archive)
    (out/"external_release_evidence_gate_summary.json").write_text(json.dumps(summary,indent=2),encoding="utf-8")
    return summary


def self_test() -> None:
    schema={"type":"object","required":["a"],"properties":{"a":{"type":"integer"}}}
    ok,errors=validate_schema({"a":1},schema)
    assert ok and not errors
    ok,errors=validate_schema({"a":"x"},schema)
    assert not ok and errors

    records=[
        {"id":f"H{i:02d}","hold":"x","closure_status":"OPEN","submission_status":"OPEN_NO_EXTERNAL_EVIDENCE"}
        for i in range(1,9)
    ]
    gate=approval_gate(records)
    assert gate["gate_state"]=="HOLD_NO_EXTERNAL_RELEASE_EVIDENCE_INGESTED"
    assert gate["open_release_holds"]==8
    assert gate["release_allowed"] is False
    print("PHOENIX_4_41_EXTERNAL_RELEASE_EVIDENCE_GATE_SELF_TEST=PASS")


def verify_output(path: Path) -> None:
    d=load_json(path)
    out=path.parent
    assert d["status"]==STATUS
    assert d["construction_release"]=="LOCKED"
    assert d["release_allowed"] is False
    required=[
        "evidence_intake_register.json",
        "evidence_intake_register.csv",
        "final_approval_gate.json",
        "release_evidence_submission.schema.json",
        "PHOENIX_ANIJSTRAAT_616_EXTERNAL_RELEASE_EVIDENCE_APPROVAL_GATE.docx",
        "PHOENIX_ANIJSTRAAT_616_EXTERNAL_RELEASE_EVIDENCE_APPROVAL_GATE.pdf",
        "evidence_manifest.json",
        "PHOENIX_ANIJSTRAAT_616_EXTERNAL_RELEASE_EVIDENCE_GATE_PACKAGE.zip",
    ]
    for rel in required:
        p=out/rel
        assert p.exists() and p.stat().st_size>0, rel

    from pypdf import PdfReader
    assert len(PdfReader(str(out/"PHOENIX_ANIJSTRAAT_616_EXTERNAL_RELEASE_EVIDENCE_APPROVAL_GATE.pdf")).pages)>=2
    from docx import Document
    assert len(Document(str(out/"PHOENIX_ANIJSTRAAT_616_EXTERNAL_RELEASE_EVIDENCE_APPROVAL_GATE.docx")).paragraphs)>20

    print("PHOENIX_4_41_EXTERNAL_RELEASE_EVIDENCE_GATE_OUTPUT_VERIFY=PASS")


def main()->int:
    ap=argparse.ArgumentParser()
    ap.add_argument("--bootstrap-inbox",type=Path)
    ap.add_argument("--schema",type=Path)
    ap.add_argument("--final-package-root",type=Path)
    ap.add_argument("--evidence-root",type=Path)
    ap.add_argument("--source-pdf",type=Path)
    ap.add_argument("--output",type=Path)
    ap.add_argument("--self-test",action="store_true")
    ap.add_argument("--verify-output",type=Path)
    args=ap.parse_args()

    if args.self_test:
        self_test(); return 0
    if args.bootstrap_inbox:
        if not args.schema:
            ap.error("--schema required with --bootstrap-inbox")
        bootstrap_inbox(args.bootstrap_inbox,args.schema)
        print(f"EVIDENCE_INBOX_BOOTSTRAP=PASS")
        print(f"EVIDENCE_ROOT={args.bootstrap_inbox}")
        return 0
    if args.verify_output:
        verify_output(args.verify_output); return 0

    if not all([args.final_package_root,args.evidence_root,args.schema,args.source_pdf,args.output]):
        ap.error("--final-package-root --evidence-root --schema --source-pdf --output required")

    summary=run_gate(args.final_package_root,args.evidence_root,args.schema,args.source_pdf,args.output)
    print(f"PROJECT_ID={summary['project_id']}")
    print(f"INGESTED_SUBMISSIONS={summary['ingested_submission_count']}")
    print(f"REJECTED_SUBMISSIONS={summary['rejected_submission_count']}")
    print(f"OPEN_RELEASE_HOLDS={summary['open_release_holds']}")
    print(f"FINAL_APPROVAL_GATE_STATE={summary['gate_state']}")
    print("FOR_CONSTRUCTION_RELEASE=LOCKED")
    print("RELEASE_ALLOWED=FALSE")
    print("EXTERNAL_RELEASE_EVIDENCE_INGESTION_GATE=PASS")
    print(f"ARCHIVE_SHA256={summary['archive_sha256']}")
    print(f"NEXT_STAGE={summary['next_stage']}")
    print(f"OUTPUT={args.output}")
    return 0


if __name__=="__main__":
    raise SystemExit(main())
